"""
Excel report of everything the autopilot has done: who was scouted, who was emailed, follow ups, replies, opt outs
and bounces. GitHub Actions builds it every morning and sends it to your Telegram, so you never have to run it.

  python report_excel.py                     build outreach_report.xlsx from the database in this folder
  python report_excel.py --telegram          build it and send it to your Telegram chat
  gh workflow run report.yml                 (from your PC) get a fresh report on Telegram right now

Sheets: Summary, Email record (every email sent: date, address, business, service, country, city, website, phone,
the exact message, follow up, status and reply), Replies (answer these), Queue (approved, not sent yet, with the email
that will go out), All leads.
The file is never saved in the repository: it holds email addresses, and the repository is public.
"""
import argparse
import json
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import get_connection
from scout import init_scout_tables
from scrapers.direct_web import COUNTRY_NAMES
from services import SERVICES, load_mailbox_map

LOCAL_OFFSET = timedelta(hours=1)  # Lagos time, UTC+1 all year
LOCAL_LABEL = "Lagos time"
SERVICES_SHOWN = ["whatsapp_api", "receptionist", "render_3d"]
SHORT = {"whatsapp_api": "WhatsApp", "receptionist": "Receptionist", "render_3d": "3D renders"}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {
    "Replied": PatternFill("solid", fgColor="C6EFCE"),
    "Opted out": PatternFill("solid", fgColor="E7E6E6"),
    "Bounced": PatternFill("solid", fgColor="FFC7CE"),
    "Waiting for reply": PatternFill("solid", fgColor="FFF2CC"),
}


def local_time(value):
    """SQLite UTC text -> datetime in Lagos time (Excel shows it as a real date)."""
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "")) + LOCAL_OFFSET
    except ValueError:
        return None


def friendly_status(r) -> str:
    s, sent = r["status"], r["sent_at"]
    if s == "replied":
        return "Replied"
    if s == "lost":
        return "Opted out"
    if s == "contacted":
        return "Waiting for reply"
    if s == "skipped" and sent:
        return "Bounced"
    if s == "approved":
        return "In queue"
    if s == "shortlisted":
        return "Shortlisted (not approved)"
    if s == "no_email":
        return "No email found"
    if s == "skipped":
        return "Skipped: " + (r["skip_reason"] or "not a fit")
    return s or ""


def emails_of(r):
    try:
        return json.loads(r["emails_json"] or "[]")
    except ValueError:
        return []


def col(r, name):
    """Column value, or None for databases made before that column existed."""
    return r[name] if name in r.keys() else None


def add_sheet(wb, title, headers, rows, widths, wrap_cols=(), status_col=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, datetime):
                c.number_format = "dd mmm yyyy hh:mm"
            if c.column in wrap_cols:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        if status_col and row[status_col - 1].value in STATUS_FILL:
            row[status_col - 1].fill = STATUS_FILL[row[status_col - 1].value]
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build(path: str = "outreach_report.xlsx") -> dict:
    init_scout_tables()
    mailboxes = load_mailbox_map()
    with get_connection() as conn:
        leads = conn.execute("SELECT * FROM scout_leads ORDER BY created_at DESC").fetchall()
        bounces_7d = conn.execute(
            "SELECT COUNT(*) FROM blacklist WHERE reason='bounce' AND created_at >= datetime('now','-7 days')"
        ).fetchone()[0]

    now_local = datetime.now(timezone.utc).replace(tzinfo=None) + LOCAL_OFFSET
    today = now_local.date()
    week_ago = now_local - timedelta(days=7)
    sent = [r for r in leads if r["sent_at"]]
    replied = [r for r in sent if r["status"] == "replied"]

    # ---------------------------------------------------------------- Summary
    wb = Workbook()
    summary_rows, totals = [], {}
    for sid in SERVICES_SHOWN:
        mine = [r for r in leads if r["service_id"] == sid]
        s_sent = [r for r in mine if r["sent_at"]]
        n_sent = len(s_sent)
        n_rep = sum(1 for r in s_sent if r["status"] == "replied")
        row = {
            "Service": SHORT[sid],
            "Mailbox": mailboxes.get(sid, "not set"),
            "Sent today": sum(1 for r in s_sent if (local_time(r["sent_at"]) or now_local).date() == today),
            "Sent last 7 days": sum(1 for r in s_sent if (local_time(r["sent_at"]) or now_local) >= week_ago),
            "Sent in total": n_sent,
            "Follow ups sent": sum(1 for r in s_sent if col(r, "followup_sent_at")),
            "Replies": n_rep,
            "Reply rate": f"{(n_rep / n_sent * 100):.1f}%" if n_sent else "",
            "Opted out": sum(1 for r in s_sent if r["status"] == "lost"),
            "Bounced": sum(1 for r in s_sent if r["status"] == "skipped"),
            "Waiting to be sent": sum(1 for r in mine if r["status"] == "approved"),
            "Scouted in total": len(mine),
        }
        totals[sid] = row
        summary_rows.append(list(row.values()))
    headers = list(next(iter(totals.values())).keys())
    ws = add_sheet(wb, "Summary", headers, summary_rows, [16, 30, 11, 13, 12, 13, 9, 10, 10, 9, 14, 13], first=True)
    ws.append([])
    ws.append([f"Report made {now_local:%d %b %Y %H:%M} ({LOCAL_LABEL}). Bounces in the last 7 days: {bounces_7d}."])
    ws.append([])
    # emails sent per day, last 14 days
    ws.append(["Emails sent per day"] + [SHORT[s] for s in SERVICES_SHOWN] + ["Total"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for d in range(13, -1, -1):
        day = today - timedelta(days=d)
        counts = [sum(1 for r in sent if r["service_id"] == s and (local_time(r["sent_at"]) or now_local).date() == day)
                  for s in SERVICES_SHOWN]
        ws.append([day.strftime("%a %d %b")] + counts + [sum(counts)])

    # ---------------------------------------------------------------- Replies
    reply_rows = []
    for r in sorted([r for r in sent if r["replied_at"]], key=lambda r: r["replied_at"], reverse=True):
        reply_rows.append([
            local_time(r["replied_at"]), friendly_status(r), SERVICES[r["service_id"]].label if r["service_id"] in SERVICES else r["service_id"],
            r["name"], col(r, "sent_to") or (emails_of(r)[:1] or [""])[0], r["country"], r["reply_snippet"] or "",
            r["sender_inbox"], r["website"],
        ])
    add_sheet(wb, "Replies", ["Replied (" + LOCAL_LABEL + ")", "Status", "Service", "Business", "Their email", "Country",
                              "What they wrote", "Reply from your mailbox", "Website"],
              reply_rows, [18, 16, 26, 28, 30, 8, 60, 28, 32], wrap_cols=(7,), status_col=2)

    # ---------------------------------------------------------------- Email record: one row per email sent
    # Everything about each email in one place: when, who, where, which service, the exact message, the follow up,
    # and what happened (waiting, replied, opted out, bounced).
    record_rows = []
    for r in sorted(sent, key=lambda r: r["sent_at"], reverse=True):
        city = (r["city"] or "").strip()
        if city.lower() in ("general", (COUNTRY_NAMES.get(r["country"] or "", "") or "").lower()):
            city = ""
        addr = (col(r, "sent_to") or (emails_of(r)[:1] or [""])[0]).lower()
        record_rows.append([
            local_time(r["sent_at"]),
            addr,
            r["name"],
            SHORT.get(r["service_id"], r["service_id"]),
            COUNTRY_NAMES.get(r["country"] or "", r["country"]),
            city,
            r["website"],
            r["phone"] or "",
            r["owner_name"] or "",
            r["keyword"] or "",
            r["sender_inbox"],
            r["sent_subject"] or "",
            col(r, "sent_body") or "(sent before message saving was added)",
            friendly_status(r),
            local_time(col(r, "followup_sent_at")),
            col(r, "followup_body") or ("" if not col(r, "followup_sent_at") else "(sent before message saving was added)"),
            local_time(r["replied_at"]),
            r["reply_snippet"] or "",
            ", ".join(e for e in emails_of(r) if e.lower() != addr),
        ])
    add_sheet(wb, "Email record", [
        "Date sent (" + LOCAL_LABEL + ")", "Email address", "Business", "Service", "Country", "City / address",
        "Website", "Phone", "Contact name", "Found by searching", "Sent from", "Subject", "Message sent", "Status",
        "Follow up sent", "Follow up message", "Replied on", "Their reply", "Other emails found"],
        record_rows, [18, 32, 28, 13, 16, 28, 32, 18, 18, 22, 28, 34, 70, 18, 18, 60, 18, 50, 30],
        wrap_cols=(13, 16, 18), status_col=14)

    # ---------------------------------------------------------------- Queue
    queue_rows = []
    from outreach import build_message
    for r in sorted([r for r in leads if r["status"] == "approved"], key=lambda r: (r["service_id"], -(r["score"] or 0))):
        ev = " | ".join(json.loads(r["evidence_json"] or "[]")) if r["evidence_json"] else ""
        try:
            subj, body = build_message(r["service_id"], dict(r))
            preview = f"Subject: {subj}\n\n{body}"
        except Exception:
            preview = ""
        queue_rows.append([SERVICES[r["service_id"]].label if r["service_id"] in SERVICES else r["service_id"], r["name"],
                           r["website"], r["country"], ", ".join(emails_of(r)), r["score"], ev, local_time(r["created_at"]),
                           preview])
    add_sheet(wb, "Queue", ["Service", "Business", "Website", "Country", "Emails found", "Score", "Why it fits", "Scouted",
                            "Email that will be sent"],
              queue_rows, [26, 28, 32, 8, 40, 7, 60, 18, 70], wrap_cols=(7, 9))

    # ---------------------------------------------------------------- All leads
    all_rows = []
    for r in leads:
        all_rows.append([local_time(r["created_at"]), SERVICES[r["service_id"]].label if r["service_id"] in SERVICES else r["service_id"],
                         r["name"], r["website"], r["country"], r["keyword"], r["phone"], ", ".join(emails_of(r)), r["owner_name"],
                         r["score"], friendly_status(r)])
    add_sheet(wb, "All leads", ["Scouted (" + LOCAL_LABEL + ")", "Service", "Business", "Website", "Country", "Search",
                                "Phone", "Emails found", "Owner", "Score", "Status"],
              all_rows, [18, 26, 28, 32, 8, 22, 18, 40, 20, 7, 30], status_col=11)

    order = ["Summary", "Email record", "Replies", "Queue", "All leads"]
    wb._sheets.sort(key=lambda ws: order.index(ws.title) if ws.title in order else 99)
    wb.active = 1  # opens on the Email record
    wb.save(path)
    return {"path": path, "sent": len(sent), "replied": len(replied), "leads": len(leads), "totals": totals,
            "queue": len(queue_rows), "sent_today": sum(t["Sent today"] for t in totals.values())}


def caption(info: dict) -> str:
    import tg
    lines = [f"<b>Outreach report</b> ({datetime.now(timezone.utc) + LOCAL_OFFSET:%d %b %H:%M})"]
    for sid, t in info["totals"].items():
        lines.append(f"{tg.esc(sid)}: sent today {t['Sent today']}, total {t['Sent in total']}, replies {t['Replies']}, "
                     f"queue {t['Waiting to be sent']}")
    lines.append(f"All: {info['sent']} emailed, {info['replied']} replied, {info['leads']} scouted")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="outreach_report.xlsx")
    ap.add_argument("--telegram", action="store_true", help="send the file to your Telegram chat")
    a = ap.parse_args()
    info = build(a.out)
    print(f"Saved {a.out}: {info['sent']} emails sent, {info['replied']} replies, {info['leads']} leads, {info['queue']} in the queue")
    if a.telegram:
        import tg
        ok = tg.send_document(a.out, caption(info))
        print("Sent to Telegram." if ok else "Telegram did not accept the file (check TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID).")


if __name__ == "__main__":
    main()
